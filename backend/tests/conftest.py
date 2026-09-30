import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.deps import get_analyzer
from app.config import Settings
from app.main import app
from app.schemas.llm import AtsCheckResult, RecommendationsResult, ScoringResult, SkillsResult
from app.services.analyzer import ResumeAnalyzer

RESUME_TEXT = (
    "Jane Doe — jane@example.com — Software Engineer with 5 years of experience "
    "building Python and FastAPI services. BSc Computer Science. Skills: Python, SQL, Docker."
)
JOB_TEXT = (
    "We are hiring a Backend Engineer with 3+ years of Python, FastAPI, PostgreSQL "
    "and Kubernetes experience. Bachelor's degree required."
)

ATS_PAYLOAD = {
    "ats_score": "82%",
    "detected_skills": ["Python", "python", "SQL", "Docker"],
    "detected_sections": ["Contact Info", "Experience", "Education", "Skills"],
    "missing_sections": ["Summary"],
    "formatting_issues": ["Inconsistent date formats"],
    "strengths": ["Clear section headers"],
    "recommendations": [
        {"priority": "low", "category": "formatting", "title": "Unify dates",
         "action": "Use one date format.", "reason": "Parsers expect consistency."},
        {"priority": "CRITICAL", "category": "content", "title": "Add a summary",
         "action": "Open with a 2-line summary.", "reason": "Recruiters skim."},
        "not-an-object",
    ],
    "summary": "Solid, parseable resume.",
}
SKILLS_PAYLOAD = {
    "matched_skills": ["Python", "FastAPI", "SQL"],
    "missing_skills": ["Kubernetes"],
    "keyword_coverage_percent": 70,
}
SCORING_PAYLOAD = {
    "match_percent": 76,
    "ats_score": 81,
    "experience": {"years_found": 5, "years_required": "3+", "verdict": "Above"},
    "education": {"found": "BSc Computer Science", "required": "Bachelor's", "match": True},
}
RECS_PAYLOAD = {
    "recommendations": [
        {"priority": "high", "category": "skills", "title": "Show Kubernetes exposure",
         "action": "Mention any container orchestration work.", "reason": "The JD requires it."},
    ],
    "learning_path": [{"skill": "Kubernetes", "reason": "Required by the role."}],
    "summary": "Good fit with one gap.",
}


class FakeLLM:
    """Stands in for LLMClient: returns canned payloads keyed by schema."""

    configured = True

    def __init__(self, payloads=None, failures=None):
        self.payloads = {
            AtsCheckResult: ATS_PAYLOAD,
            SkillsResult: SKILLS_PAYLOAD,
            ScoringResult: SCORING_PAYLOAD,
            RecommendationsResult: RECS_PAYLOAD,
            **(payloads or {}),
        }
        self.failures = failures or {}
        self.prompts = []

    def complete(self, model, prompt, schema, *, max_tokens=None, json_mode=False):
        self.prompts.append((schema, prompt))
        if schema in self.failures:
            raise self.failures[schema]
        return schema.model_validate(self.payloads[schema])


@pytest.fixture
def settings():
    return Settings(_env_file=None, groq_api_key="test-key")


@pytest.fixture
def fake_llm():
    return FakeLLM()


@pytest.fixture
def client(fake_llm, settings):
    app.dependency_overrides[get_analyzer] = lambda: ResumeAnalyzer(fake_llm, settings)
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def resume_file():
    return {"resume": ("resume.txt", RESUME_TEXT.encode(), "text/plain")}


@pytest.fixture
def docx_bytes():
    doc = Document()
    doc.add_paragraph("Jane Doe — Software Engineer")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, FastAPI, Docker, PostgreSQL"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
