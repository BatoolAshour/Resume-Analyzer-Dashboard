"""Extract and clean text from uploaded resume / job description files."""

import io
import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf  # PyMuPDF
from docx import Document

from app.config import ALLOWED_EXTENSIONS
from app.errors import EmptyDocument, FileTooLarge, InvalidFileType, UnreadableFile

# Anything shorter than this is not a usable resume / job description.
MIN_TEXT_CHARS = 40


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    text: str
    word_count: int
    truncated: bool


def _extract_pdf(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        return "\n".join(page.get_text() for page in pdf)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Resumes are often laid out in tables, which doc.paragraphs skips.
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                # Merged cells are returned once per column they span.
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    if b"\x00" in data:
        raise UnreadableFile("This does not look like a plain-text file.")
    return data.decode("utf-8", errors="ignore")


_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".txt": _extract_txt}


def extract_text(filename: str, data: bytes) -> str:
    """Pull raw text out of a PDF, DOCX, or TXT upload."""
    extension = Path(filename or "").suffix.lower()
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise InvalidFileType(
            f"Unsupported file type '{extension or 'unknown'}'. "
            f"Upload one of: {', '.join(e.lstrip('.').upper() for e in ALLOWED_EXTENSIONS)}."
        )
    try:
        return extractor(data)
    except UnreadableFile:
        raise
    except Exception as exc:
        raise UnreadableFile(
            f"Could not read '{filename}'. The file may be corrupted or password-protected."
        ) from exc


def clean_text(text: str) -> str:
    """Tidy whitespace but keep line breaks.

    Headings, bullets and contact details sit on their own lines; flattening
    everything into one line makes a well-structured resume look unstructured.
    """
    lines = (re.sub(r"[^\S\n]+", " ", line).strip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def parse_document(
    filename: str,
    data: bytes,
    *,
    max_bytes: int,
    max_chars: int,
    label: str = "resume",
) -> ParsedDocument:
    """Validate an upload and turn it into cleaned text ready for the analyzer."""
    if not data:
        raise EmptyDocument(f"The uploaded {label} file is empty.")
    if len(data) > max_bytes:
        raise FileTooLarge(
            f"The {label} file is too large. The limit is {max_bytes // (1024 * 1024)} MB."
        )

    text = clean_text(extract_text(filename, data))
    if len(text) < MIN_TEXT_CHARS:
        raise EmptyDocument(
            f"No readable text was found in the {label}. "
            "Scanned or image-only files are not supported — upload a text-based file."
        )
    return parse_text(text, filename=filename, max_chars=max_chars)


def parse_text(text: str, *, filename: str, max_chars: int) -> ParsedDocument:
    text = clean_text(text)
    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]
    return ParsedDocument(
        filename=filename,
        text=text,
        word_count=len(text.split()),
        truncated=truncated,
    )
